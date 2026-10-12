"""Profiler for flexible sales file imports with arbitrary column names."""

from __future__ import annotations

import csv
import io
import json
import re
import unicodedata
from dataclasses import asdict, dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable, Mapping


SUPPORTED_SALES_EXTENSIONS = {".csv", ".tsv", ".json", ".jsonl", ".xlsx"}
_IDENTIFIER_PATTERN = re.compile(r"[^a-z0-9]+")

# Common column name aliases mapped to canonical field names
SEMANTIC_ALIASES: dict[str, tuple[str, ...]] = {
    "sales_date": (
        "order_date", "invoice_date", "sale_date", "sales_date", "transaction_date",
        "purchase_date", "date", "created_at",
    ),
    "order_id": (
        "order_id", "order_number", "invoice_id", "invoice_number", "transaction_id",
        "transaction_number", "bill_no", "bill_number", "receipt_number",
    ),
    "customer": (
        "customer_id", "customer_name", "customer", "client_id", "client", "client_name",
        "buyer", "buyer_name", "account", "account_name",
    ),
    "product": (
        "product_id", "product_name", "product", "item_name", "item", "sku", "item_code",
        "description",
    ),
    "category": ("product_category", "category", "product_group", "department", "segment"),
    "region": ("sales_region", "region", "territory", "zone", "state", "city", "country"),
    "quantity": ("quantity", "qty", "units_sold", "units", "volume", "item_quantity"),
    "unit_price": ("unit_price", "selling_price", "sale_price", "price", "rate", "unit_rate"),
    "discount": ("discount_amount", "discount_value", "discount", "discounts"),
    "revenue": (
        "net_revenue", "net_sales", "sales_amount", "sale_amount", "revenue", "sales",
        "line_total", "total_amount", "grand_total", "amount", "subtotal", "total",
        "total_sales", "total_price", "order_total", "order_amount", "item_total", "line_amount",
    ),
    "cost": ("cost_amount", "total_cost", "unit_cost", "cost", "cogs"),
    "profit": ("profit_amount", "gross_profit", "net_profit", "profit", "margin_amount"),
    "status": ("order_status", "sales_status", "payment_status", "status"),
    "currency": ("currency_code", "currency", "currency_symbol", "curr", "curr_code", "cur", "currency_name"),
}


class SalesImportError(ValueError):
    """Raised when a file cannot safely be treated as a sales export."""


@dataclass(frozen=True)
class SalesImportProfile:
    """A persisted explanation of how a flexible file was interpreted."""

    source_files: list[str]
    source_row_counts: dict[str, int]
    row_count: int
    accepted_row_count: int
    rejected_row_count: int
    columns: list[str]
    column_types: dict[str, str]
    semantic_mapping: dict[str, str]
    warnings: list[str]

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


@dataclass(frozen=True)
class FlexibleSalesImport:
    """Wide source rows and a canonical representation of the same sales data."""

    rows: list[dict[str, Any]]
    canonical_rows: list[dict[str, Any]]
    profile: SalesImportProfile


def is_supported_sales_file(path: Path) -> bool:
    return path.is_file() and path.suffix.lower() in SUPPORTED_SALES_EXTENSIONS


def load_flexible_sales_files(paths: Iterable[Path]) -> FlexibleSalesImport:
    """Read CSV/TSV/JSON/JSONL/XLSX exports and infer a conservative sales map.

    Values remain parameterized Python values; no source text is ever used as
    SQL.  A caller can save the returned profile beside the immutable raw copy
    and show its warnings to the person who uploaded the file.
    """

    source_paths = [Path(path).resolve() for path in paths]
    if not source_paths:
        raise SalesImportError("Upload at least one CSV, TSV, JSON, JSONL, or XLSX sales file.")

    unsupported = [str(path) for path in source_paths if not is_supported_sales_file(path)]
    if unsupported:
        raise SalesImportError("Unsupported sales file type: " + ", ".join(unsupported))

    source_rows: list[tuple[Path, int, dict[str, Any]]] = []
    source_row_counts: dict[str, int] = {}
    warnings: list[str] = []
    for path in source_paths:
        rows = _read_source_rows(path)
        if not rows:
            warnings.append(f"{path.name} contains no data rows and was skipped.")
            continue
        source_row_counts[path.name] = len(rows)
        for row_number, row in enumerate(rows, start=2):
            source_rows.append((path, row_number, _normalise_record(row)))

    if not source_rows:
        raise SalesImportError("No usable sales records were found in the supplied file(s).")

    columns = _ordered_columns(row for _, _, row in source_rows)
    semantic_mapping, mapping_warnings = _infer_semantic_mapping(columns)
    warnings.extend(mapping_warnings)
    _validate_sales_signals(semantic_mapping, columns)

    column_types = _infer_column_types([row for _, _, row in source_rows], columns)
    typed_rows: list[dict[str, Any]] = []
    canonical_rows: list[dict[str, Any]] = []
    rows_without_measure = 0
    mapped_columns = set(semantic_mapping.values())

    for source_path, row_number, row in source_rows:
        typed = {column: _coerce_value(row.get(column), column_types[column]) for column in columns}
        typed_rows.append(typed)
        canonical = _canonical_row(
            typed,
            source_file=source_path.name,
            source_row=row_number,
            mapping=semantic_mapping,
            mapped_columns=mapped_columns,
        )
        if canonical["revenue"] is None:
            rows_without_measure += 1
        canonical_rows.append(canonical)

    if rows_without_measure:
        warnings.append(
            f"{rows_without_measure:,} row(s) have no usable revenue value; they remain explorable but are excluded from revenue totals."
        )
    if "sales_date" not in semantic_mapping:
        warnings.append("No sales date was detected, so time-series questions may not be answerable.")
    if "order_id" not in semantic_mapping:
        warnings.append("No order or invoice ID was detected; order-count metrics use source rows instead.")

    profile = SalesImportProfile(
        source_files=[path.name for path in source_paths],
        source_row_counts=source_row_counts,
        row_count=len(source_rows),
        accepted_row_count=len(canonical_rows),
        rejected_row_count=0,
        columns=columns,
        column_types=column_types,
        semantic_mapping=semantic_mapping,
        warnings=warnings,
    )
    return FlexibleSalesImport(rows=typed_rows, canonical_rows=canonical_rows, profile=profile)


def _read_source_rows(path: Path) -> list[dict[str, Any]]:
    suffix = path.suffix.lower()
    if suffix in {".csv", ".tsv"}:
        return _read_delimited(path, delimiter="\t" if suffix == ".tsv" else None)
    if suffix in {".json", ".jsonl"}:
        return _read_json(path)
    if suffix == ".xlsx":
        return _read_xlsx(path)
    raise SalesImportError(f"Unsupported sales file type: {path.name}")


def _read_delimited(path: Path, delimiter: str | None) -> list[dict[str, Any]]:
    try:
        text = path.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError as error:
        raise SalesImportError(f"{path.name} must be UTF-8 encoded.") from error
    if not text.strip():
        return []
    if delimiter is None:
        try:
            dialect = csv.Sniffer().sniff(text[:16_384], delimiters=",;|\t")
            delimiter = dialect.delimiter
        except csv.Error:
            delimiter = ","
    reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
    if not reader.fieldnames or not any((field or "").strip() for field in reader.fieldnames):
        raise SalesImportError(f"{path.name} has no usable header row.")
    return [dict(row) for row in reader if any(value not in (None, "") for value in row.values())]


def _read_json(path: Path) -> list[dict[str, Any]]:
    try:
        if path.suffix.lower() == ".jsonl":
            parsed: Any = [json.loads(line) for line in path.read_text(encoding="utf-8-sig").splitlines() if line.strip()]
        else:
            parsed = json.loads(path.read_text(encoding="utf-8-sig"))
    except (UnicodeDecodeError, json.JSONDecodeError) as error:
        raise SalesImportError(f"{path.name} is not valid UTF-8 JSON.") from error
    if isinstance(parsed, list):
        return [dict(item) for item in parsed if isinstance(item, Mapping)]
    if isinstance(parsed, Mapping):
        list_values = [value for value in parsed.values() if isinstance(value, list) and all(isinstance(item, Mapping) for item in value)]
        if len(list_values) == 1:
            return [dict(item) for item in list_values[0]]
        return [dict(parsed)]
    raise SalesImportError(f"{path.name} must contain an object, an array of objects, or JSON Lines.")


def _read_xlsx(path: Path) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as error:  # pragma: no cover - exercised in deployments without extras
        raise SalesImportError(
            "XLSX import needs the 'openpyxl' dependency. Run `pip install -e .` after pulling this release."
        ) from error
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as error:  # noqa: BLE001
        raise SalesImportError(f"Could not read spreadsheet {path.name}.") from error

    try:
        sheet = workbook.active
        if sheet is None and workbook.worksheets:
            sheet = workbook.worksheets[0]
        if sheet is None:
            raise SalesImportError(f"{path.name} contains no active worksheet.")
        if not hasattr(sheet, "iter_rows"):
            raise SalesImportError(f"{path.name} active worksheet does not support row iteration.")

        iterator = sheet.iter_rows(values_only=True)
        header_values: Any = None
        for values in iterator:
            if values and any(v is not None and str(v).strip() for v in values):
                header_values = values
                break
        if not header_values:
            return []
        headers = [str(value).strip() if value is not None else "" for value in header_values]
        if not any(headers):
            raise SalesImportError(f"{path.name} has no usable header row.")
        rows: list[dict[str, Any]] = []
        for values in iterator:
            if not values or not any(value not in (None, "") for value in values):
                continue
            rows.append({headers[index]: value for index, value in enumerate(values) if index < len(headers)})
        return rows
    except SalesImportError:
        raise
    except Exception as error:  # noqa: BLE001
        raise SalesImportError(f"Could not read spreadsheet {path.name}.") from error
    finally:
        workbook.close()


def _normalise_identifier(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode("ascii").lower()
    text = _IDENTIFIER_PATTERN.sub("_", text).strip("_")
    if not text:
        text = "column"
    if text[0].isdigit():
        text = f"field_{text}"
    return text[:63]


def _normalise_record(record: Mapping[str, Any]) -> dict[str, Any]:
    flattened = _flatten_record(record)
    result: dict[str, Any] = {}
    used: set[str] = set()
    for original, value in flattened.items():
        base = _normalise_identifier(original)
        candidate = base
        counter = 2
        while candidate in used:
            candidate = f"{base}_{counter}"
            counter += 1
        used.add(candidate)
        result[candidate] = value
    return result


def _flatten_record(record: Mapping[str, Any], prefix: str = "") -> dict[str, Any]:
    flattened: dict[str, Any] = {}
    for key, value in record.items():
        name = f"{prefix}_{key}" if prefix else key
        if isinstance(value, Mapping):
            flattened.update(_flatten_record(value, name))
        elif isinstance(value, list):
            flattened[name] = json.dumps(value, ensure_ascii=False, default=str)
        else:
            flattened[name] = value
    return flattened


def _ordered_columns(rows: Iterable[dict[str, Any]]) -> list[str]:
    seen: set[str] = set()
    columns: list[str] = []
    for row in rows:
        for column in row:
            if column not in seen:
                seen.add(column)
                columns.append(column)
    return columns


def _infer_semantic_mapping(columns: list[str]) -> tuple[dict[str, str], list[str]]:
    mapping: dict[str, str] = {}
    warnings: list[str] = []
    column_set = set(columns)
    for semantic, aliases in SEMANTIC_ALIASES.items():
        matches = [alias for alias in aliases if alias in column_set]
        if matches:
            mapping[semantic] = matches[0]
            if len(matches) > 1:
                warnings.append(
                    f"Mapped {semantic!r} to {matches[0]!r}; also found {', '.join(matches[1:])}. Review this mapping if their meanings differ."
                )
    return mapping, warnings


def _validate_sales_signals(mapping: Mapping[str, str], columns: list[str]) -> None:
    has_revenue = "revenue" in mapping
    has_derived_revenue = "quantity" in mapping and "unit_price" in mapping
    identity_signals = sum(field in mapping for field in ("order_id", "product", "sales_date", "customer"))
    if not has_revenue and not has_derived_revenue and identity_signals < 2:
        preview = ", ".join(columns[:12])
        raise SalesImportError(
            "This does not look like a sales export. Include a revenue/amount field, or quantity and unit price, "
            f"plus sales identifiers. Detected columns: {preview or 'none'}."
        )


def _infer_column_types(rows: list[dict[str, Any]], columns: list[str]) -> dict[str, str]:
    types: dict[str, str] = {}
    for column in columns:
        values = [row.get(column) for row in rows[:2_000] if not _is_blank(row.get(column))]
        if not values:
            types[column] = "VARCHAR"
            continue
        parsed_dates = [_parse_date(value) for value in values]
        if sum(value is not None for value in parsed_dates) / len(values) >= 0.95:
            types[column] = "DATE"
            continue
        parsed_numbers = [_parse_number(value) for value in values]
        if sum(value is not None for value in parsed_numbers) / len(values) >= 0.95:
            numeric_values = [value for value in parsed_numbers if value is not None]
            types[column] = "BIGINT" if all(value == value.to_integral_value() for value in numeric_values) else "DOUBLE"
            continue
        types[column] = "VARCHAR"
    return types


def _coerce_value(value: Any, duckdb_type: str) -> Any:
    if _is_blank(value):
        return None
    if duckdb_type == "DATE":
        parsed = _parse_date(value)
        return parsed.isoformat() if parsed else None
    if duckdb_type in {"DOUBLE", "BIGINT"}:
        parsed = _parse_number(value)
        if parsed is None:
            return None
        return int(parsed) if duckdb_type == "BIGINT" else float(parsed)
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, default=str)
    return str(value).strip()


def _canonical_row(
    row: Mapping[str, Any],
    *,
    source_file: str,
    source_row: int,
    mapping: Mapping[str, str],
    mapped_columns: set[str],
) -> dict[str, Any]:
    def mapped_value(field: str) -> Any:
        """Return the primary mapped value or a compatible alias in this row.

        A folder can legitimately contain exports from two sales systems.  For
        example, one file may call the measure ``sales_amount`` while another
        calls it ``revenue``.  The profile presents one deterministic primary
        mapping, but canonicalization must not silently discard values from
        the other accepted alias.
        """

        candidates: list[str] = []
        primary = mapping.get(field)
        if primary:
            candidates.append(primary)
        candidates.extend(alias for alias in SEMANTIC_ALIASES[field] if alias not in candidates)
        for column in candidates:
            value = row.get(column)
            if not _is_blank(value):
                return value
        return None

    def text(field: str) -> str | None:
        value = mapped_value(field)
        return str(value).strip() if not _is_blank(value) else None

    def number(field: str) -> float | None:
        value = mapped_value(field)
        parsed = _parse_number(value)
        return float(parsed) if parsed is not None else None

    sales_date = _parse_date(mapped_value("sales_date"))
    quantity = number("quantity")
    unit_price = number("unit_price")
    discount = number("discount") or 0.0
    revenue = number("revenue")
    if revenue is None:
        if quantity is not None and unit_price is not None:
            revenue = max((quantity * unit_price) - discount, 0.0)
        elif unit_price is not None:
            revenue = max(unit_price - discount, 0.0)
    elif quantity is None and unit_price is not None and unit_price > 0:
        quantity = max(round(revenue / unit_price, 2), 1.0)

    # Detect currency from column, or fallback to currency symbols in raw measures
    currency_val = text("currency")
    if not currency_val:
        for field in ("revenue", "unit_price", "price", "amount", "total"):
            raw_field_val = mapped_value(field) if field in SEMANTIC_ALIASES else row.get(field)
            detected = _detect_currency_symbol(str(raw_field_val or ""))
            if detected:
                currency_val = detected
                break
    else:
        norm_sym = _detect_currency_symbol(currency_val)
        if norm_sym:
            currency_val = norm_sym

    attributes = {key: value for key, value in row.items() if key not in mapped_columns and value is not None}
    customer_str = text("customer") or "Unknown"
    product_str = text("product") or "Unknown Product"
    return {
        "source_file": source_file,
        "source_row": source_row,
        "sales_date": sales_date.isoformat() if sales_date else None,
        "order_id": text("order_id"),
        "customer": customer_str,
        "customer_id": customer_str,
        "customer_name": customer_str,
        "product": product_str,
        "product_name": product_str,
        "category": text("category") or "Unspecified",
        "region": text("region") or "Unspecified",
        "quantity": quantity,
        "unit_price": unit_price,
        "discount": discount if "discount" in mapping else None,
        "revenue": revenue,
        "net_revenue": revenue,
        "cost": number("cost"),
        "profit": number("profit"),
        "status": text("status") or "completed",
        "currency": currency_val,
        "attributes_json": json.dumps(attributes, ensure_ascii=False, default=str) if attributes else None,
    }


def _detect_currency_symbol(text: str) -> str | None:
    if not text:
        return None
    if "$" in text or re.search(r"\b(usd|dollars?)\b", text, re.I):
        return "$"
    if "₹" in text or re.search(r"\b(inr|rs\.?|rupees?)\b", text, re.I):
        return "₹"
    if "€" in text or re.search(r"\b(eur|euros?)\b", text, re.I):
        return "€"
    if "£" in text or re.search(r"\b(gbp|pounds?)\b", text, re.I):
        return "£"
    return None


def _parse_number(value: Any) -> Decimal | None:
    if _is_blank(value):
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float, Decimal)):
        try:
            return Decimal(str(value))
        except InvalidOperation:
            return None
    text = str(value).strip()
    negative = text.startswith("(") and text.endswith(")")
    text = text.strip("()")
    # Strip currency symbols and common currency codes before the identifier check
    # so values like "$1,200", "USD 1200", "₹5,000", "INR 5000", or "Rs. 5000" are parsed as numbers.
    stripped_for_alpha_check = re.sub(
        r"[\s₹$€£¥(),.%-]|(?i:\b(usd|inr|eur|gbp|cad|aud|rs\.?|rupees?|dollars?)\b\.?)",
        "",
        text,
    )
    # Values that mix alphabetic characters with digits (e.g. "INV-1",
    # "ORD-002", "SKU-100") are identifiers, not numbers.
    if re.search(r"[a-zA-Z]", stripped_for_alpha_check) and re.search(r"[0-9]", stripped_for_alpha_check):
        return None
    # Remove currency tokens and symbols from text so "Rs. 5000" does not leave an unwanted dot
    clean_text = re.sub(r"(?i)\b(usd|inr|eur|gbp|cad|aud|rs\.?|rupees?|dollars?)\b\.?", "", text)
    clean_text = re.sub(r"[\s₹$€£¥]", "", clean_text)
    clean_text = re.sub(r"[^0-9,.-]", "", clean_text).replace(",", "")
    if not clean_text or clean_text in {"-", ".", "-."}:
        return None
    try:
        number = Decimal(clean_text)
        return -number if negative else number
    except InvalidOperation:
        return None


def _parse_date(value: Any) -> date | None:
    if _is_blank(value):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).date()
    except ValueError:
        pass
    for pattern in ("%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y", "%m-%d-%Y", "%d %b %Y", "%d %B %Y"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    return None


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())
